"""Evidencia sobre PostgreSQL efímero: nunca usa ni modifica pedidos comerciales."""
import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.infrastructure.commercial_exports import currency_format
from apps.core.models import Pedido, ItemPedido, Tienda, ConsentimientoComercial

ORDERS = '/api/tienda/pedidos/historial/'
SALES = '/api/tienda/ventas/historial/'
CONTACTS = '/api/tienda/contactos/'
EXPORTS = ['/api/tienda/informes/xlsx/', '/api/tienda/pedidos/exportar/', '/api/tienda/ventas/exportar/', '/api/tienda/contactos/exportar/']
FILTER = {'fecha_desde': '2026-10-01', 'fecha_hasta': '2026-10-31'}


class CommercialHistoryTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.owner = users.objects.create_user(username='commercial-owner')
        self.other = users.objects.create_user(username='commercial-other')
        self.store = Tienda.objects.create(nombre='Tienda QA XLSX', id_usuario_propietario=self.owner)
        self.api = APIClient(); self.api.force_authenticate(self.owner)
        self.paid = self.order('QA-PAID', 'pagado', '18500.50', ' Buyer@Example.com ', discount='2500', shipping='1000', price='10000.25', quantity=2)
        self.delivered = self.order('QA-DELIVERED', 'entregado', '900.25', 'buyer@example.COM', discount='100', price='1000.25')
        self.pending = self.order('QA-PENDING', 'pendiente', '7000', 'pending@example.com', price='7000', delivery='Chilexpress')
        self.cancelled = self.order('QA-CANCELLED', 'cancelado', '9000', 'cancelled@example.com', price='9000')

    def order(self, identifier, state='pagado', amount='1000', email='test@example.com', *, discount='0', shipping='0', price=None, quantity=1, delivery='Retiro', instant='2026-10-10T15:00:00+00:00'):
        obj = Pedido.objects.create(tienda=self.store, identificador=identifier, estado=state, monto_total=amount,
                correo_contacto=email, nombre_contacto='Cliente QA', contacto={'nombre': 'Cliente QA'}, descuento=discount, costo_envio=shipping,
                medio_pago='Transferencia', metodo_pago='Transferencia', entrega={'metodo': delivery, 'retiro': {'seleccion': {'fecha': '2026-11-02', 'hora': '10:07'}}})
        Pedido.objects.filter(pk=obj.pk).update(fecha_creacion=datetime.fromisoformat(instant)); obj.refresh_from_db()
        if price:
            ItemPedido.objects.create(id_pedido=obj, nombre_producto='Producto histórico QA', sku='SKU-QA', cantidad=quantity, precio_unitario=price,
                                      atributos={'Color': 'Azul'}, atributos_variante=[{'nombre': 'Color', 'valor': 'Azul'}])
        return obj

    def get_data(self, url, params=None):
        response = self.api.get(url, params or {})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response['Cache-Control'], 'no-store')
        return response.json()

    def binary(self, url, params=None):
        response = self.api.get(url, params or {})
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        self.assertEqual(response['Cache-Control'], 'no-store')
        content = b''.join(response.streaming_content)
        # El wrapper del cliente de Django cierra el stream al consumirlo.
        return content

    def workbook(self, url, params=None):
        content = self.binary(url, params)
        self.assertTrue(content.startswith(b'PK'))
        return load_workbook(BytesIO(content), data_only=False)

    def sheet_rows(self, book, sheet):
        return list(book[sheet].values)[3:]

    def test_01_dashboard_real_xlsx_summary_reconciles_with_json(self):
        book = self.workbook(EXPORTS[0], {'anio': 2026, 'mes': 10})
        self.assertEqual(book.sheetnames, ['Resumen', 'Comparación mensual', 'Estados', 'Pedidos', 'Líneas', 'Productos vendidos'])
        summary = dict(self.sheet_rows(book, 'Resumen'))
        self.assertEqual(Decimal(str(summary['Importe con pago confirmado'])), Decimal('19400.75'))
        self.assertEqual(Decimal(str(summary['Descuentos de pedidos confirmados'])), Decimal('2600'))
        report = self.get_data('/api/tienda/informes/', {'anio': 2026, 'mes': 10})['data']
        self.assertEqual(Decimal(report['total_ventas']), Decimal(str(summary['Importe con pago confirmado'])))
        book.close()

    def test_02_month_date_filter_uses_chile_and_exclusive_end(self):
        self.order('SEPT-UTC', instant='2026-10-01T02:30:00+00:00')
        self.order('NOVEMBER', instant='2026-11-01T03:00:00+00:00')
        book = self.workbook(EXPORTS[0], {'anio': 2026, 'mes': 10})
        self.assertEqual(len(self.sheet_rows(book, 'Pedidos')), 4); book.close()
        orders = self.get_data(ORDERS, FILTER)
        self.assertEqual(orders['paginacion']['total'], 4)

    def test_03_orders_detail_search_and_pagination(self):
        page = self.get_data(ORDERS, {'q': 'QA-', 'tamano': 1, 'pagina': 2})
        self.assertEqual(len(page['data']), 1); self.assertEqual(page['paginacion']['total'], 4)
        row = self.get_data(f'/api/tienda/pedidos/{self.paid.pk}/detalle/')['data']
        self.assertEqual(row['items'][0]['sku'], 'SKU-QA')
        self.assertEqual(row['items'][0]['cantidad'], 2)
        self.assertEqual(row['entrega']['retiro']['seleccion']['hora'], '10:07')
        self.assertEqual(row['pago_estado'], 'confirmado')

    def test_04_delivery_payment_and_group_filters(self):
        self.assertEqual(self.get_data(ORDERS, {'entrega': 'envio'})['paginacion']['total'], 1)
        self.assertEqual(self.get_data(ORDERS, {'entrega': 'retiro', 'pago': 'confirmado'})['paginacion']['total'], 2)
        self.assertEqual(self.get_data(ORDERS, {'estado': 'confirmed', 'metodo_pago': 'transferencia'})['paginacion']['total'], 2)
        self.assertEqual(self.get_data(ORDERS, {'estado': 'dispatch'})['paginacion']['total'], 2)

    def test_05_orders_export_includes_all_pages_same_filters(self):
        book = self.workbook(EXPORTS[1], {**FILTER, 'q': 'QA-', 'tamano': 1, 'pagina': 2})
        self.assertEqual(len(self.sheet_rows(book, 'Pedidos')), 4)
        statuses = {r[3] for r in self.sheet_rows(book, 'Pedidos')}
        self.assertEqual(statuses, {'confirmado', 'cancelado', 'pendiente'})
        self.assertTrue(all(r[11] == 0 for r in self.sheet_rows(book, 'Pedidos') if r[3] != 'confirmado')); book.close()

    def test_06_sales_excludes_cancelled_pending_and_preserves_lines(self):
        data = self.get_data(SALES, FILTER)['data']
        self.assertEqual(data['resumen']['cantidad_ventas'], 2)
        self.assertEqual(data['resumen']['total_vendido'], '19400.75')
        self.assertEqual(sum(len(o['items']) for o in data['ventas']), 2)
        self.assertTrue(all(o['pago_estado'] == 'confirmado' for o in data['ventas']))

    def test_07_sales_export_totals_discounts_and_products(self):
        book = self.workbook(EXPORTS[2], FILTER)
        summary = dict(self.sheet_rows(book, 'Resumen'))
        self.assertEqual(Decimal(str(summary['Bruto productos (total + descuento − envío)'])), Decimal('21000.75'))
        self.assertEqual(Decimal(str(summary['Neto productos confirmados'])), Decimal('18400.75'))
        self.assertEqual(sum(r[2] for r in self.sheet_rows(book, 'Productos vendidos')), 3)
        self.assertEqual(len(self.sheet_rows(book, 'Pedidos')), 2)
        self.assertEqual(self.sheet_rows(book, 'Por período')[0][1], 2); book.close()

    def test_08_no_invented_line_discount(self):
        data = self.get_data(SALES)['data']
        self.assertIsNone(data['ventas'][0]['items'][0]['descuento_linea'])
        book = self.workbook(EXPORTS[2]); self.assertTrue(all(r[7] is None for r in self.sheet_rows(book, 'Líneas'))); book.close()

    def test_09_clp_usd_and_native_excel_dates_and_numbers(self):
        self.assertIn('CLP', currency_format()); self.assertIn('USD', currency_format('USD'))
        book = self.workbook(EXPORTS[1])
        self.assertIsInstance(book['Pedidos']['B4'].value, datetime)
        self.assertEqual(book['Pedidos']['B4'].number_format, 'dd/mm/yyyy hh:mm')
        self.assertIsInstance(book['Pedidos']['K4'].value, (int, float))
        self.assertIn('CLP', book['Pedidos']['K4'].number_format)
        self.assertEqual(book['Pedidos'].freeze_panes, 'A4'); self.assertIsNotNone(book['Pedidos'].auto_filter.ref); book.close()

    def test_10_contacts_case_trim_dedup_and_count(self):
        rows = self.get_data(CONTACTS)['data']
        buyer = next(r for r in rows if r['email'] == 'buyer@example.com')
        self.assertEqual(buyer['cantidad_pedidos'], 2)
        self.assertEqual(len(rows), 3)

    def test_11_invalid_empty_email_discarded_and_legacy_source(self):
        self.order('INVALID', email='not-email')
        legacy = self.order('LEGACY', email=''); legacy.contacto = {'email': ' Legacy@Example.com '}; legacy.save(update_fields=['contacto'])
        blank = self.order('BLANK', email=''); blank.contacto = {'nombre': 'Sin email'}; blank.save(update_fields=['contacto'])
        rows = self.get_data(CONTACTS)['data']
        self.assertEqual({r['email'] for r in rows}, {'buyer@example.com', 'pending@example.com', 'cancelled@example.com', 'legacy@example.com'})

    def test_12_private_endpoints_unauthenticated_other_and_inactive(self):
        for user, expected in ((None, 401), (self.other, 403)):
            self.api.force_authenticate(user)
            for url in [ORDERS, SALES, CONTACTS, *EXPORTS, f'/api/tienda/pedidos/{self.paid.pk}/detalle/']:
                with self.subTest(url=url, expected=expected): self.assertEqual(self.api.get(url).status_code, expected)
        self.owner.is_active = False; self.owner.save(update_fields=['is_active']); self.api.force_authenticate(self.owner)
        self.assertEqual(self.api.get(CONTACTS).status_code, 403)

    def test_13_forged_tenant_and_unassigned_orders_never_leak(self):
        orphan = self.order('ORPHAN', email='orphan@example.com'); orphan.tienda = None; orphan.save(update_fields=['tienda'])
        for url in [ORDERS, SALES, CONTACTS, *EXPORTS]:
            self.assertEqual(self.api.get(url, {'tienda_id': self.store.pk + 99}).status_code, 400)
        self.assertEqual(self.get_data(ORDERS)['paginacion']['total'], 4)
        self.assertNotIn('orphan@example.com', [r['email'] for r in self.get_data(CONTACTS)['data']])
        self.assertEqual(self.api.get(f'/api/tienda/pedidos/{orphan.pk}/detalle/').status_code, 404)

    def consent(self, email, state='autorizado'):
        return ConsentimientoComercial.objects.create(tienda=self.store, email=email, estado=state,
                    version_texto='qa-v1', referencia_evidencia='fixture-explicit-opt-in')

    def test_14_marketing_explicit_opt_in_only_and_withdrawal(self):
        record = self.consent(' BUYER@example.COM ')
        self.consent('cancelled@example.com', 'retirado')
        book = self.workbook(EXPORTS[3], {'uso': 'marketing'})
        self.assertEqual([r[0] for r in self.sheet_rows(book, 'Contactos')], ['buyer@example.com']); book.close()
        record.estado = 'retirado'; record.save(update_fields=['estado'])
        book = self.workbook(EXPORTS[3], {'uso': 'marketing'}); self.assertEqual(self.sheet_rows(book, 'Contactos'), []); book.close()

    def test_15_purchase_or_cookie_is_not_consent(self):
        self.paid.contacto = {'consent': True, 'cookies': 'accepted'}; self.paid.save(update_fields=['contacto'])
        self.assertTrue(all(r['consentimiento'] == 'sin_registro' for r in self.get_data(CONTACTS)['data']))
        book = self.workbook(EXPORTS[3], {'uso': 'marketing'}); self.assertEqual(self.sheet_rows(book, 'Contactos'), []); book.close()
        with self.assertRaises(ValidationError):
            ConsentimientoComercial.objects.create(tienda=self.store, email='no-evidence@example.com', estado='autorizado')

    def test_16_contacts_filters_and_exports_csv_xlsx(self):
        self.consent('buyer@example.com')
        rows = self.get_data(CONTACTS, {'q': 'BUYER', 'consentimiento': 'autorizado', **FILTER})
        self.assertEqual(rows['paginacion']['total'], 1)
        content = self.binary(EXPORTS[3], {'formato': 'csv', 'q': 'buyer'})
        self.assertTrue(content.startswith(b'\xef\xbb\xbf'))
        parsed = list(csv.reader(StringIO(content.decode('utf-8-sig')))); self.assertEqual(parsed[1][0], 'buyer@example.com')
        self.assertEqual(parsed[1][2], '10/10/2026 12:00')
        book = self.workbook(EXPORTS[3], {'q': 'buyer'}); self.assertEqual(len(self.sheet_rows(book, 'Contactos')), 1); book.close()

    def test_17_formula_injection_stays_text_in_xlsx_and_csv(self):
        line = self.paid.itempedido_set.get(); line.nombre_producto = '=HYPERLINK("https://example.test","x")'; line.sku = '@SUM(1)'; line.save()
        self.order('=1+1', email='=formula@example.com')
        book = self.workbook(EXPORTS[1])
        for sheet in book:
            for row in sheet:
                self.assertTrue(all(cell.data_type != 'f' for cell in row))
        self.assertTrue(any(r[2].startswith('=HYPERLINK') for r in self.sheet_rows(book, 'Líneas'))); book.close()
        content = self.binary(EXPORTS[3], {'formato': 'csv', 'q': '=formula'})
        self.assertIn("'=formula@example.com", content.decode('utf-8-sig'))

    def test_18_invalid_filters_are_errors_not_silent_exports(self):
        for params in ({'fecha_desde': '31/10/2026'}, {'fecha_desde': '2026-11-02', 'fecha_hasta': '2026-10-01'},
                       {'tamano': '101'}, {'pago': 'unknown'}, {'entrega': 'teleport'}, {'private': 'x'}):
            for url in (ORDERS, SALES, EXPORTS[1], EXPORTS[2]):
                with self.subTest(params=params, url=url): self.assertEqual(self.api.get(url, params).status_code, 400)
        self.assertEqual(self.api.get(ORDERS + '?q=one&q=two').status_code, 400)
        self.assertEqual(self.api.get(EXPORTS[3], {'uso': 'broadcast'}).status_code, 400)

    def test_19_empty_workbooks_and_legacy_orders_without_lines(self):
        legacy = self.order('OLD-NO-LINES')
        self.assertEqual(self.get_data(SALES)['data']['importes']['pedidos_sin_lineas'], 1)
        book = self.workbook(EXPORTS[2], {'q': 'DOES-NOT-EXIST'})
        self.assertEqual(self.sheet_rows(book, 'Pedidos'), []); self.assertEqual(self.sheet_rows(book, 'Líneas'), []); book.close()

    def test_20_export_limits_are_explicit(self):
        with patch('apps.core.infrastructure.commercial.MAX_ORDERS', 1):
            response = self.api.get(EXPORTS[1]); self.assertEqual(response.status_code, 400)
            self.assertEqual(response.data['error']['codigo'], 'EXPORTACION_DEMASIADO_GRANDE')

    def test_21_queries_are_bounded_and_no_n_plus_one(self):
        for i in range(40): self.order(f'SCALE-{i}', price='1000')
        with CaptureQueriesContext(connection) as queries:
            response = self.api.get(ORDERS, {'tamano': 20})
        self.assertEqual(response.status_code, 200); self.assertLessEqual(len(queries), 6)
        with CaptureQueriesContext(connection) as queries:
            response = self.api.get(CONTACTS)
        self.assertEqual(response.status_code, 200); self.assertLessEqual(len(queries), 4)

    def test_22_multiple_product_matches_do_not_duplicate_order_totals(self):
        ItemPedido.objects.create(id_pedido=self.paid, nombre_producto='Producto histórico QA extra', sku='SKU-QA2', cantidad=1, precio_unitario=0)
        rows = self.get_data(SALES, {'q': 'histórico QA'})['data']
        self.assertEqual(rows['resumen']['cantidad_ventas'], 2)
        self.assertEqual(rows['resumen']['total_vendido'], '19400.75')
        self.assertEqual(rows['importes']['descuentos'], '2600.00')
        self.assertEqual(rows['importes']['bruto_productos'], '21000.75')
        self.assertEqual(sum(r['cantidad'] for r in rows['productos']), 4)
        self.assertEqual(sum(Decimal(r['total_vendido']) for r in rows['serie']), Decimal('19400.75'))

    def test_23_pdf_compatibility(self):
        response = self.api.get('/api/tienda/informes/pdf/', {'anio': 2026, 'mes': 10})
        self.assertEqual(response.status_code, 200); self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_24_contact_list_has_no_grant_consent_or_bulk_email_action(self):
        self.assertEqual(self.api.post(CONTACTS, {'email': 'buyer@example.com', 'consentimiento': 'autorizado'}, format='json').status_code, 405)
        self.assertFalse(ConsentimientoComercial.objects.exists())

    def test_25_many_orders_and_whitespace_still_one_contact(self):
        for i in range(10): self.order(f'DUP-{i}', email='\t\nBUYER@example.com\u00a0')
        self.consent('buyer@example.com')
        data = self.get_data(CONTACTS, {'q': 'buyer', 'consentimiento': 'autorizado'})
        self.assertEqual(data['paginacion']['total'], 1)
        self.assertEqual(data['data'][0]['cantidad_pedidos'], 12)
        book = self.workbook(EXPORTS[3], {'uso': 'marketing'})
        self.assertEqual(len(self.sheet_rows(book, 'Contactos')), 1); book.close()

    def test_26_store_query_scope_and_consent_are_independent(self):
        foreign = Tienda.objects.create(nombre='Otro tenant QA', id_usuario_propietario=self.other)
        order = self.order('FOREIGN', email='foreign@example.com'); order.tienda = foreign; order.save(update_fields=['tienda'])
        ConsentimientoComercial.objects.create(tienda=foreign, email='buyer@example.com', estado='autorizado', version_texto='qa-v1', referencia_evidencia='other-tenant')
        # La instalación con dos tiendas falla de forma cerrada.
        self.assertEqual(self.api.get(CONTACTS).status_code, 503)
        # Defensa de consulta independiente del resolver: incluso autorizada una tienda, no mezcla fuentes.
        with patch('apps.core.presentation.commercial_exports.resolver_tienda', return_value=self.store):
            rows = self.get_data(CONTACTS)['data']
            self.assertNotIn('foreign@example.com', [row['email'] for row in rows])
            book = self.workbook(EXPORTS[3], {'uso': 'marketing'})
            self.assertEqual(self.sheet_rows(book, 'Contactos'), []); book.close()
        with patch('apps.core.presentation.historial.resolver_tienda', return_value=self.store):
            self.assertNotIn('FOREIGN', [row['identificador'] for row in self.get_data(ORDERS)['data']])
            self.assertEqual(self.api.get(f'/api/tienda/pedidos/{order.pk}/detalle/').status_code, 404)
