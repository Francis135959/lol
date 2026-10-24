"""Exportadores del servidor: XLSX Open XML y CSV, sin datos bancarios/personales ajenos al fin."""
import csv
import re
from datetime import datetime, date
from decimal import Decimal
from tempfile import SpooledTemporaryFile

from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.core.infrastructure.commercial import amounts, payment_status, paid_summary, product_totals, local_date, export_check
from apps.core.infrastructure.reports import NOTA_FECHAS
from apps.core.models import Pedido

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def currency_format(currency='CLP'):
    if currency not in ('CLP', 'USD'):
        raise ValueError('Moneda no admitida')
    return f'#,##0.00 "{currency}";[Red]-#,##0.00 "{currency}"'


def safe_text(value):
    return ILLEGAL_XML.sub('', str(value))[:32767]


def safe_csv(value):
    value = safe_text(value)
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')) else value


class ExportBook:
    def __init__(self, store, label):
        self.book = Workbook(write_only=True)
        self.book.properties.creator = 'OpenMind'
        self.book.properties.title = safe_text(label)
        self.store = safe_text(store.nombre)
        self.label = safe_text(label)

    def cell(self, sheet, value, *, header=False, money=False):
        cell = WriteOnlyCell(sheet, value=safe_text(value) if isinstance(value, str) else value)
        if isinstance(value, str):
            # Nunca interpretar nombres, SKU o email como fórmulas, hipervínculos o errores Excel.
            cell.data_type = 's'
        cell.font = Font(name='Calibri', size=11, color='FFFFFF' if header else '172B4D', bold=header)
        cell.alignment = Alignment(vertical='top', wrap_text=isinstance(value, str))
        if header:
            cell.fill = PatternFill('solid', fgColor='184E63')
        elif money or isinstance(value, Decimal):
            cell.number_format = currency_format()
        elif isinstance(value, datetime):
            cell.number_format = 'dd/mm/yyyy hh:mm'
        elif isinstance(value, date):
            cell.number_format = 'dd/mm/yyyy'
        elif isinstance(value, int):
            cell.number_format = '#,##0'
        return cell

    def sheet(self, name, headers, rows, *, currency_columns=(), note=''):
        sheet = self.book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'A4'
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 42 if index == 1 else max(19, min(38, len(header) + 4))
        if name == 'Resumen':
            sheet.column_dimensions['A'].width = 58
            sheet.column_dimensions['B'].width = 88
        sheet.row_dimensions[1].height = 27
        sheet.row_dimensions[2].height = 44
        sheet.append([self.cell(sheet, f'{self.store} · {self.label}', header=True)])
        sheet.append([self.cell(sheet, note or 'Importes CLP · fechas en Chile · exportación comercial no tributaria.')])
        sheet.append([self.cell(sheet, h, header=True) for h in headers])
        count = 3
        for row in rows:
            sheet.append([self.cell(sheet, v, money=i in currency_columns) for i, v in enumerate(row)])
            count += 1
        sheet.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{count}'
        sheet.print_title_rows = '1:3'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    def response(self, filename):
        output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
        try:
            self.book.save(output)
            output.seek(0)
            response = FileResponse(output, as_attachment=True, filename=filename, content_type=XLSX_MIME)
            response['Cache-Control'] = 'no-store'
            response['X-Content-Type-Options'] = 'nosniff'
            return response
        except BaseException:
            output.close()
            raise


def export_orders(store, query, label, filename, *, report=None, series=None, filters=None):
    export_check(query)
    book = ExportBook(store, label)
    summary = paid_summary(query)
    summary_rows = [
        ('Pedidos del período', query.count()), ('Pedidos confirmados', summary['pedidos_confirmados']),
        ('Bruto productos (total + descuento − envío)', Decimal(summary['bruto_productos'])),
        ('Descuentos de pedidos confirmados', Decimal(summary['descuentos'])),
        ('Neto productos confirmados', Decimal(summary['neto_productos'])),
        ('Envío confirmado', Decimal(summary['envio'])), ('Importe con pago confirmado', Decimal(summary['importe_confirmado'])),
        ('Pedidos confirmados sin líneas históricas', summary['pedidos_sin_lineas']),
        ('Criterio de fecha', NOTA_FECHAS), ('Descuentos por producto', summary['nota_descuentos']),
        ('Fuente de confirmación', 'Estado actual del pedido. No acredita liquidación bancaria ni calcula utilidades.'),
        ('Filtros aplicados', '; '.join(f'{key}: {value}' for key, value in (filters or {}).items() if key not in ('pagina', 'tamano', 'tienda_id', 'id_tienda')) or 'Todo el histórico'),
    ]
    book.sheet('Resumen', ['Concepto', 'Valor'], summary_rows)
    if report:
        book.sheet('Comparación mensual', ['Mes', 'Pedidos', 'Confirmados', 'Importe confirmado CLP'],
                   ((date(report['anio'], row['mes'], 1), row['pedidos'], row['pagados'], Decimal(row['ventas'])) for row in report['meses']))
    if series is not None:
        book.sheet('Por período', ['Período', 'Pedidos confirmados', 'Importe confirmado CLP'],
                   ((r['periodo'], r['cantidad_ventas'], Decimal(r['total_vendido'])) for r in series))
    from django.db.models import Count
    states = query.order_by().values('estado').annotate(cantidad=Count('pk')).order_by('estado')
    book.sheet('Estados', ['Estado', 'Cantidad de pedidos'], ((r['estado'], r['cantidad']) for r in states))

    def order_rows():
        for order in query.iterator(chunk_size=500):
            a = amounts(order)
            yield (order.identificador or f'#{order.pk}', local_date(order.fecha_creacion), order.get_estado_display(),
                   payment_status(order), order.metodo_pago or order.medio_pago, (order.entrega or {}).get('metodo', 'Sin registro'),
                   a['bruto_productos'], a['descuento'], a['neto_productos'], a['envio'], a['total'], a['confirmado'])
    book.sheet('Pedidos', ['Pedido', 'Registro en Chile', 'Estado', 'Estado de pago', 'Método de pago', 'Entrega',
                          'Bruto productos CLP', 'Descuento CLP', 'Neto productos CLP', 'Envío CLP', 'Total pedido CLP', 'Confirmado CLP'], order_rows())

    def line_rows():
        for order in query.prefetch_related('itempedido_set').iterator(chunk_size=200):
            for line in order.itempedido_set.all():
                # Descuento de línea no disponible: no se reparte ni se repite el descuento del pedido.
                attributes = line.atributos_variante or line.atributos
                import json
                yield (order.identificador or f'#{order.pk}', local_date(order.fecha_creacion), line.nombre_producto,
                       line.sku, json.dumps(attributes, ensure_ascii=False) if attributes else '', line.cantidad,
                       line.precio_unitario, None, line.precio_unitario * line.cantidad, payment_status(order),
                       order.metodo_pago or order.medio_pago)
    book.sheet('Líneas', ['Pedido', 'Registro en Chile', 'Producto', 'SKU', 'Variante', 'Cantidad', 'Precio unitario CLP',
                          'Descuento línea (no registrado)', 'Bruto línea CLP', 'Estado de pago', 'Método de pago'],
               line_rows(), note='Descuentos y total neto se consultan en Pedidos. Las líneas muestran su precio histórico; no equivalen a ingresos netos.')
    book.sheet('Productos vendidos', ['Producto', 'SKU', 'Cantidad confirmada', 'Bruto líneas CLP'],
               ((r['nombre_producto'], r['sku'], r['unidades'], r['bruto_lineas']) for r in product_totals(query)),
               note='Solo pedidos con pago confirmado. Importe de líneas antes del descuento del pedido; no incluye envío.')
    return book.response(filename)


def export_contacts(store, rows, *, csv_format=False, marketing=False):
    purpose = 'marketing-autorizado' if marketing else 'operacional'
    filename = f'contactos-{purpose}-{timezone.now():%Y%m%d}'
    headers = ['Email', 'Cantidad de pedidos', 'Último pedido en Chile', 'Consentimiento comercial', 'Fecha de consentimiento en Chile']
    def values():
        for row in rows:
            yield (row['email'], row['cantidad_pedidos'], local_date(datetime.fromisoformat(row['ultimo_pedido'])),
                   row['consentimiento'], local_date(datetime.fromisoformat(row['fecha_consentimiento'])) if row['fecha_consentimiento'] else None)
    if csv_format:
        import io
        output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
        text = io.TextIOWrapper(output, encoding='utf-8-sig', newline='')
        writer = csv.writer(text)
        writer.writerow(headers)
        for row in values():
            writer.writerow([safe_csv(v.strftime('%d/%m/%Y %H:%M') if isinstance(v, datetime) else v) if v is not None else '' for v in row])
        text.flush(); text.detach(); output.seek(0)
        response = FileResponse(output, as_attachment=True, filename=filename + '.csv', content_type='text/csv; charset=utf-8')
        response['Cache-Control'] = 'no-store'
        response['X-Content-Type-Options'] = 'nosniff'
        return response
    book = ExportBook(store, 'Contactos · ' + purpose)
    book.sheet('Contactos', headers, values(), note='Un correo por contacto. Comprar no implica autorizar marketing. ' + ('Solo autorizaciones expresas registradas.' if marketing else 'Uso operativo de pedidos; no es una lista comercial autorizada.'))
    return book.response(filename + '.xlsx')
